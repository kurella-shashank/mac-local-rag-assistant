
import pytest

from mac_local_rag.ingestion.reader import read_document
from docx import Document as WordDocument
from reportlab.pdfgen import canvas

def test_txt_extraction(tmp_path):
    file = tmp_path / "sample.txt"
    file.write_text("AWS Glue processes S3 data.")

    result = read_document(file)

    assert result.content == "AWS Glue processes S3 data."
    assert result.file_type == ".txt"


def test_markdown_extraction(tmp_path):
    file = tmp_path / "notes.md"
    file.write_text("# RAG Architecture")

    result = read_document(file)

    assert result.content == "# RAG Architecture"


def test_unsupported_extension(tmp_path):
    file = tmp_path / "image.png"
    file.write_bytes(b"sample")

    with pytest.raises(ValueError, match="Unsupported"):
        read_document(file)


def test_empty_file(tmp_path):
    file = tmp_path / "empty.txt"
    file.write_text("")

    with pytest.raises(ValueError, match="empty"):
        read_document(file)


def test_invalid_encoding(tmp_path):
    file = tmp_path / "invalid.txt"
    file.write_bytes(b"\xff\xfe")

    with pytest.raises(UnicodeError):
        read_document(file)


def test_symlink_rejected(tmp_path):
    original = tmp_path / "original.txt"
    original.write_text("Private document")

    link = tmp_path / "link.txt"
    link.symlink_to(original)

    with pytest.raises(ValueError, match="Symbolic"):
        read_document(link)

        
def test_pdf_extraction(tmp_path):
    path = tmp_path / "sample.pdf"

    pdf = canvas.Canvas(str(path))
    pdf.drawString(72, 750, "AWS Glue processes S3 data.")
    pdf.save()

    result = read_document(path)

    assert result.file_type == ".pdf"
    assert "AWS Glue processes S3 data" in result.content


def test_docx_extraction(tmp_path):
    path = tmp_path / "sample.docx"

    doc = WordDocument()
    doc.add_paragraph("Qdrant stores embeddings.")
    doc.save(path)

    result = read_document(path)

    assert result.file_type == ".docx"
    assert "Qdrant stores embeddings." in result.content


def test_docx_table_extraction(tmp_path):
    path = tmp_path / "table.docx"

    doc = WordDocument()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Database"
    table.cell(0, 1).text = "Qdrant"
    doc.save(path)

    result = read_document(path)

    assert "Database | Qdrant" in result.content


def test_pdf_without_text(tmp_path):
    path = tmp_path / "empty.pdf"

    pdf = canvas.Canvas(str(path))
    pdf.showPage()
    pdf.save()

    with pytest.raises(ValueError, match="No extractable"):
        read_document(path)
